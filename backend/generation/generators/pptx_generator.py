"""PowerPoint presentation (.pptx) generator for technical briefings."""

from __future__ import annotations

import io
from datetime import datetime, timezone
from typing import TYPE_CHECKING

if TYPE_CHECKING:
    from generation.document_models import DocumentOutline, GeneratedSection


class PptxGenerator:
    """Builds authentic PowerPoint (.pptx) slide decks from outlines and sections."""

    def build_pptx_bytes(
        self,
        outline: DocumentOutline,
        sections: list[GeneratedSection],
    ) -> bytes:
        try:
            from pptx import Presentation
            from pptx.util import Inches, Pt
            from pptx.dml.color import RGBColor
        except ImportError as exc:
            raise ImportError(
                "python-pptx is required to generate PowerPoint decks. Please run: pip install python-pptx"
            ) from exc

        prs = Presentation()

        # Slide 1: Title Slide
        title_slide_layout = prs.slide_layouts[0]
        slide = prs.slides.add_slide(title_slide_layout)
        title = slide.shapes.title
        subtitle = slide.placeholders[1]

        title.text = outline.title
        now_str = datetime.now(timezone.utc).strftime("%Y-%m-%d")
        subtitle.text = f"{outline.intended_audience}\nOmniOps Sovereign Industrial Intelligence • {now_str}"

        # Subsequent Content Slides
        bullet_slide_layout = prs.slide_layouts[1]

        for sec in sections:
            slide = prs.slides.add_slide(bullet_slide_layout)
            slide.shapes.title.text = sec.title

            body_shape = slide.placeholders[1]
            tf = body_shape.text_frame
            tf.word_wrap = True

            # First bullet
            first_bullet = True
            for line in sec.content_markdown.splitlines():
                clean_line = line.strip()
                if not clean_line or clean_line.startswith("### "):
                    continue
                if first_bullet:
                    tf.text = clean_line[:120]
                    first_bullet = False
                else:
                    p = tf.add_paragraph()
                    p.text = clean_line[:120]
                    p.level = 0

            # Add callout bullets
            for callout in sec.callouts:
                c_type = callout.get("type", "WARNING").upper()
                c_text = callout.get("text", "")
                p = tf.add_paragraph()
                p.text = f"[{c_type}] {c_text[:110]}"
                p.level = 1
                p.font.bold = True
                p.font.color.rgb = RGBColor(0xB9, 0x1C, 0x1C) if c_type in {"WARNING", "DANGER"} else RGBColor(0xD9, 0x77, 0x06)

            # If section contains a table, add a structured table on the slide if space permits
            if sec.tables:
                tbl_data = sec.tables[0]
                headers = tbl_data.get("headers", [])
                rows = tbl_data.get("rows", [])
                if headers and rows:
                    rows_count = min(len(rows) + 1, 4)
                    cols_count = min(len(headers), 4)
                    left = Inches(1.0)
                    top = Inches(4.5)
                    width = Inches(8.0)
                    height = Inches(2.0)

                    table_shape = slide.shapes.add_table(rows_count, cols_count, left, top, width, height)
                    table = table_shape.table

                    # Headers
                    for c_idx in range(cols_count):
                        cell = table.cell(0, c_idx)
                        cell.text = str(headers[c_idx])
                        cell.text_frame.paragraphs[0].font.size = Pt(11)
                        cell.text_frame.paragraphs[0].font.bold = True

                    # Rows
                    for r_idx in range(1, rows_count):
                        for c_idx in range(cols_count):
                            cell = table.cell(r_idx, c_idx)
                            cell.text = str(rows[r_idx - 1][c_idx])
                            cell.text_frame.paragraphs[0].font.size = Pt(10)

        bio = io.BytesIO()
        prs.save(bio)
        return bio.getvalue()
