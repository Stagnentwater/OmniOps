"""Unit tests for WordGenerator (.docx format)."""

from __future__ import annotations

import io
import unittest

from docx import Document

from generation.document_models import DocumentOutline, GeneratedSection, OutlineSection
from generation.generators.word_generator import WordGenerator


class TestWordGenerator(unittest.TestCase):
    """Test suite for Word document creation, formatting, and structural validation."""

    def setUp(self) -> None:
        self.generator = WordGenerator()

    def test_build_docx_bytes_and_open(self) -> None:
        """Verify building .docx returns valid bytes readable by python-docx."""
        outline = DocumentOutline(
            title="Technical Work Permit: Pump P-101A",
            deliverable_format="docx",
            intended_audience="Operations & Mechanical Maintenance",
            sections=[
                OutlineSection(title="1. Scope of Work"),
                OutlineSection(title="2. Isolation Protocol"),
            ],
        )

        sections = [
            GeneratedSection(
                section_id="sec-1",
                title="1. Scope of Work",
                content_markdown="Perform mechanical seal replacement on crude pump P-101A.",
                callouts=[{"type": "WARNING", "text": "Lockout breaker MCC-4B-02."}],
            ),
            GeneratedSection(
                section_id="sec-2",
                title="2. Isolation Protocol",
                content_markdown="Verified double block and bleed valves closed.",
                tables=[{
                    "table_name": "isolation_valves",
                    "headers": ["Valve Tag", "Location", "Position", "Status"],
                    "rows": [
                        ["MOV-101A", "Pump Suction", "CLOSED", "LOCKED"],
                        ["MOV-102A", "Pump Discharge", "CLOSED", "LOCKED"],
                    ],
                }],
                citations=[{"document_id": "SOP-ISOLATION.pdf", "page_index": 5}],
            ),
        ]

        docx_bytes = self.generator.build_docx_bytes(outline, sections)
        self.assertIsInstance(docx_bytes, bytes)
        self.assertTrue(len(docx_bytes) > 1000)

        # Open and inspect the generated Word document
        doc = Document(io.BytesIO(docx_bytes))
        paragraphs_text = [p.text for p in doc.paragraphs if p.text.strip()]

        # Verify title and section headings exist
        self.assertTrue(any("Technical Work Permit: Pump P-101A" in t for t in paragraphs_text))
        self.assertTrue(any("1. Scope of Work" in t for t in paragraphs_text))
        self.assertTrue(any("2. Isolation Protocol" in t for t in paragraphs_text))

        # Verify table content
        self.assertTrue(len(doc.tables) >= 2)  # Metadata table + isolation table
        table_cells = [cell.text for row in doc.tables[-1].rows for cell in row.cells]
        self.assertIn("MOV-101A", table_cells)
        self.assertIn("LOCKED", table_cells)


if __name__ == "__main__":
    unittest.main()
