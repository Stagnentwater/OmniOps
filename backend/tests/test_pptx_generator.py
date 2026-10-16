"""Unit tests for PptxGenerator (.pptx format)."""

from __future__ import annotations

import io
import unittest

from pptx import Presentation

from generation.document_models import DocumentOutline, GeneratedSection, OutlineSection
from generation.generators.pptx_generator import PptxGenerator


class TestPptxGenerator(unittest.TestCase):
    """Test suite for PowerPoint presentation generation, layout, and slide shapes."""

    def setUp(self) -> None:
        self.generator = PptxGenerator()

    def test_build_pptx_bytes_and_inspect_slides(self) -> None:
        """Verify building .pptx creates slides readable by python-pptx."""
        outline = DocumentOutline(
            title="Refinery Turnaround 2026: Executive Briefing",
            deliverable_format="pptx",
            intended_audience="Operations Leadership & Plant Directorate",
            sections=[
                OutlineSection(title="1. Maintenance Objectives & Schedule"),
                OutlineSection(title="2. Critical Path Isolations & Safety Controls"),
            ],
        )

        sections = [
            GeneratedSection(
                section_id="sec-1",
                title="1. Maintenance Objectives & Schedule",
                content_markdown="Review shutdown timeline for Crude Unit 200.\nTarget duration is 14 days.",
                callouts=[{"type": "WARNING", "text": "Ensure hydrogen line purging is completed by Shift 1."}],
            ),
            GeneratedSection(
                section_id="sec-2",
                title="2. Critical Path Isolations & Safety Controls",
                content_markdown="Major isolation blinds required for column C-101 and reboiler E-102.",
                tables=[{
                    "table_name": "isolations",
                    "headers": ["Equipment", "Tag", "Status"],
                    "rows": [
                        ["Column", "C-101", "BLINDED"],
                        ["Reboiler", "E-102", "ISOLATED"],
                    ],
                }],
            ),
        ]

        pptx_bytes = self.generator.build_pptx_bytes(outline, sections)
        self.assertIsInstance(pptx_bytes, bytes)
        self.assertTrue(len(pptx_bytes) > 2000)

        # Inspect presentation
        prs = Presentation(io.BytesIO(pptx_bytes))
        # Total slides should be 1 (Title slide) + 2 (Content slides) = 3
        self.assertEqual(len(prs.slides), 3)

        # Title slide checks
        slide_0 = prs.slides[0]
        self.assertEqual(slide_0.shapes.title.text, "Refinery Turnaround 2026: Executive Briefing")

        # Slide 1 checks
        slide_1 = prs.slides[1]
        self.assertEqual(slide_1.shapes.title.text, "1. Maintenance Objectives & Schedule")

        # Slide 2 checks (has table)
        slide_2 = prs.slides[2]
        self.assertEqual(slide_2.shapes.title.text, "2. Critical Path Isolations & Safety Controls")
        # Verify table shape exists
        has_table = any(shape.has_table for shape in slide_2.shapes)
        self.assertTrue(has_table)


if __name__ == "__main__":
    unittest.main()
